"""Builds docs/IronPass-Presentation.pptx from the screenshots in docs/screenshots.

Run with: python docs/tools/build_presentation.py
Needs: python-pptx and pillow.
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "docs" / "IronPass-Presentation.pptx"

ACCENT = RGBColor(0xFF, 0x5A, 0x1F)
INK = RGBColor(0x16, 0x18, 0x1D)
SOFT = RGBColor(0x4A, 0x4F, 0x5A)
PAPER = RGBColor(0xF4, 0xF1, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD4, 0xC7)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def text_box(slide, x, y, w, h, text, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def bullets(slide, x, y, w, h, items, size=24, gap=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for i, item in enumerate(items):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.space_after = Pt(gap)
        dot = para.add_run()
        dot.text = "■  "
        dot.font.name = FONT
        dot.font.size = Pt(size - 8)
        dot.font.color.rgb = ACCENT
        run = para.add_run()
        run.text = item
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = INK
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, text=None, size=16, bold=False,
         color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    if text is not None:
        frame = shp.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = text if isinstance(text, list) else [text]
        for i, line_text in enumerate(lines):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = line_text
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold if i == 0 else False
            run.font.color.rgb = color
    return shp


def arrow(slide, x, y, w=0.45, h=0.35):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = ACCENT
    shp.line.fill.background()
    return shp


def picture(slide, name, x, y, w, h):
    """Place a screenshot inside the box without stretching it."""
    path = SHOTS / name
    if not path.exists():
        raise FileNotFoundError(path)
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px - 0.04), Inches(py - 0.04),
                                   Inches(pw + 0.08), Inches(ph + 0.08))
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE
    frame.shadow.inherit = False
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def new_slide(title, notes, number):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), prs.slide_height)
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()
    text_box(slide, 0.6, 0.35, 12, 0.9, title, size=36, bold=True)
    text_box(slide, 12.2, 6.95, 0.9, 0.4, str(number), size=12, color=SOFT, align=PP_ALIGN.RIGHT)
    text_box(slide, 0.6, 6.95, 6, 0.4, "IronPass  |  DevOps capstone", size=12, color=SOFT)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


# 1. title ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = INK
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.6), Inches(0.35), Inches(1.9))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
text_box(s, 1.5, 2.4, 11, 1.2, "IronPass", size=66, bold=True, color=WHITE)
text_box(s, 1.5, 3.55, 11, 1.0, "From a gym membership app to a monitored Kubernetes cluster",
         size=28, color=PAPER)
text_box(s, 1.5, 5.4, 11, 0.6, "DevOps final capstone  |  Romit", size=20, color=ACCENT, bold=True)
s.notes_slide.notes_text_frame.text = (
    "Hi, I am Romit. This is IronPass, a gym membership app. The app is small, but the point of "
    "this project is the journey. It goes from my laptop, through tests, containers and a "
    "pipeline, to a Kubernetes cluster on AWS with monitoring."
)

# 2. problem and app -----------------------------------------------------
s = new_slide("The problem and the app", (
    "A gym front desk needs to know if a member can walk in. IronPass keeps members, plans and "
    "check-ins. The API enforces one real rule. If the plan has expired, the check-in is "
    "rejected with a clear message. The frontend is React, the backend is FastAPI, and the "
    "data is in PostgreSQL."), 2)
bullets(s, 0.6, 1.4, 4.6, 5, [
    "Members, plans and check-ins",
    "Rule: expired plan means no check-in",
    "React + FastAPI + PostgreSQL",
    "12 REST endpoints, /health, /ready, /metrics",
], size=22)
picture(s, "01-app-running-localhost-3000.png", 5.4, 1.3, 7.5, 5.5)

# 3. architecture --------------------------------------------------------
s = new_slide("Architecture", (
    "This is the full picture. A developer pushes to GitHub. GitHub Actions tests, builds, scans "
    "and pushes images. Terraform builds the network and the cluster. Helm puts the app on the "
    "cluster. The Ingress is the single door. Prometheus scrapes the backend and Grafana draws "
    "the graphs."), 3)
rect(s, 0.6, 1.5, 2.0, 0.9, PAPER, text="Developer", bold=True)
arrow(s, 2.7, 1.72)
rect(s, 3.3, 1.5, 2.2, 0.9, PAPER, text="GitHub", bold=True)
arrow(s, 5.6, 1.72)
rect(s, 6.2, 1.5, 2.6, 0.9, PAPER, text="GitHub Actions", bold=True)
arrow(s, 8.9, 1.72)
rect(s, 9.5, 1.5, 2.6, 0.9, PAPER, text="GHCR images", bold=True)
rect(s, 0.6, 2.9, 12.1, 3.7, WHITE, LINE)
text_box(s, 0.8, 2.95, 8, 0.4, "AWS ap-south-1  |  VPC  |  EKS (built by Terraform)", size=14, color=SOFT, bold=True)
rect(s, 0.9, 4.0, 2.4, 1.0, PAPER, text="Ingress NGINX", bold=True)
arrow(s, 3.4, 4.3)
rect(s, 4.0, 3.5, 2.6, 0.8, ACCENT, None, text="Frontend x2", bold=True, color=WHITE)
rect(s, 4.0, 4.7, 2.6, 0.8, ACCENT, None, text="Backend x2 (HPA)", bold=True, color=WHITE)
arrow(s, 6.7, 4.9)
rect(s, 7.3, 4.7, 2.2, 0.8, PAPER, text="PostgreSQL + PVC", bold=True)
rect(s, 7.3, 5.7, 2.2, 0.7, PAPER, text="Prometheus", bold=True)
arrow(s, 9.6, 5.8)
rect(s, 10.2, 5.7, 2.2, 0.7, PAPER, text="Grafana", bold=True)
text_box(s, 0.9, 5.4, 3.0, 1.0, ["/  goes to the frontend", "/api goes to the backend"], size=14, color=SOFT)

# 4. docker --------------------------------------------------------------
s = new_slide("Run it locally with Docker", (
    "One command starts the whole stack, docker compose up --build. Postgres, the backend and "
    "the frontend all come up. The backend runs as user 10001 and nginx runs as user 101, so "
    "neither is root. The frontend image is a multi-stage build, so Node is not in the final "
    "image."), 4)
bullets(s, 0.6, 1.4, 4.4, 5, [
    "docker compose up --build",
    "3 services: frontend, backend, postgres",
    "Frontend: multi-stage, unprivileged nginx",
    "Both images run as non-root",
], size=22)
picture(s, "slides/compose.png", 5.2, 1.4, 7.7, 5.2)

# 5. tests ---------------------------------------------------------------
s = new_slide("Tests", (
    "There are seventeen pytest cases across health, plans, members, check-ins and stats. They "
    "use an in-memory SQLite database, so they are fast and safe. One test proves that an "
    "expired member cannot check in, and another proves that renewing fixes it. The pipeline "
    "runs these before any image is built."), 5)
bullets(s, 0.6, 1.4, 4.4, 5, [
    "17 tests, all passing",
    "SQLite test database",
    "Rule tested: expired plan is rejected",
    "Runs first in the pipeline",
], size=22)
picture(s, "slides/pytest.png", 5.2, 1.4, 7.7, 5.2)

# 6. git -----------------------------------------------------------------
s = new_slide("Git and GitHub", (
    "I worked in small commits, one logical change each, in build order. The messages say what "
    "changed in plain words. The gitignore keeps secrets, caches and Terraform state out of the "
    "repo."), 6)
bullets(s, 0.6, 1.4, 4.4, 5, [
    "Small commits, one change each",
    "Plain messages",
    ".gitignore: .env, state, node_modules",
    "No secrets in history",
], size=22)
picture(s, "slides/git.png", 5.2, 1.4, 7.7, 5.2)

# 7. pipeline ------------------------------------------------------------
s = new_slide("CI/CD pipeline", (
    "Every push to main runs the pipeline. First pytest and the frontend build. Then both images "
    "are built, scanned by Trivy, and only then pushed to GHCR with the commit SHA as the tag. "
    "The last job deploys with Helm. It signs in to AWS using OIDC, so there are no stored keys."), 7)
stages = ["pytest", "Build frontend", "Build images", "Trivy scan", "Push to GHCR", "Helm deploy"]
x = 0.6
for i, name in enumerate(stages):
    rect(s, x, 1.5, 1.7, 0.9, ACCENT if i in (3,) else PAPER, None if i == 3 else LINE, text=name,
         bold=True, size=14, color=WHITE if i == 3 else INK)
    if i < len(stages) - 1:
        arrow(s, x + 1.72, 1.78, 0.38, 0.32)
    x += 2.1
picture(s, "06-github-actions-green.png", 2.6, 2.7, 8.1, 4.1)

# 8. security ------------------------------------------------------------
s = new_slide("Security scan with Trivy", (
    "Trivy scans both images for high and critical problems that already have a fix. My first "
    "runs failed on purpose. The backend had seven findings and the frontend had forty one. I "
    "upgraded the libraries, removed pip from the runtime image and moved to a current nginx "
    "base. Then both scans were clean. Clean means no known fixable problem today, not safe "
    "forever."), 8)
bullets(s, 0.6, 1.4, 4.4, 5, [
    "HIGH and CRITICAL, fixable only",
    "Fails the build",
    "First runs failed: 7 and 41 findings",
    "Fixed, now clean",
], size=22)
picture(s, "08-trivy-scan.png", 5.2, 1.4, 7.7, 5.2)

# 9. terraform -----------------------------------------------------------
s = new_slide("Terraform and AWS", (
    "All the AWS infrastructure is code. A VPC with two public and two private subnets, one NAT "
    "gateway, and an EKS cluster with a managed node group. Terraform also creates the GitHub "
    "OIDC provider and a deploy role that only the main branch of this repo can assume. Every "
    "resource is tagged Project ironpass and ManagedBy terraform."), 9)
bullets(s, 0.6, 1.4, 4.6, 5, [
    "VPC: 2 public + 2 private subnets",
    "One NAT gateway",
    "EKS with managed nodes, ap-south-1",
    "GitHub OIDC role, no AWS keys",
], size=22)
picture(s, "slides/eks.png", 5.4, 1.4, 7.5, 5.2)

# 10. kubernetes and helm ------------------------------------------------
s = new_slide("Kubernetes and Helm", (
    "The Helm chart holds everything. Two Deployments with two replicas each, ClusterIP "
    "services, a PostgreSQL deployment with a volume claim, an Ingress, an HPA and a "
    "ServiceMonitor. Probes use health and ready. The ready endpoint checks the database."), 10)
bullets(s, 0.6, 1.4, 4.4, 5, [
    "One Helm chart, prod and dev values",
    "2 replicas each, ClusterIP services",
    "Probes: /health and /ready",
    "Postgres with a PVC",
], size=22)
picture(s, "slides/pods.png", 5.2, 1.4, 7.7, 5.2)

# 11. ingress and hpa ----------------------------------------------------
s = new_slide("Ingress and autoscaling", (
    "The Ingress sends slash to the frontend and slash api to the backend, on one load balancer "
    "address. For autoscaling, the HPA watches backend CPU. My load script sends steady traffic, "
    "CPU rises, and the replica count grows."), 11)
picture(s, "17-ingress-app.png", 0.6, 1.3, 6.1, 5.4)
picture(s, "slides/hpa.png", 6.9, 1.3, 6.0, 5.4)

# 12. monitoring ---------------------------------------------------------
s = new_slide("Monitoring", (
    "Prometheus scrapes the backend metrics endpoint, and the target is up. Grafana has a "
    "provisioned dashboard with request rate, latency and error rate, plus CPU and replicas. "
    "You can see them move during the load test."), 12)
picture(s, "21-grafana-dashboard.png", 0.6, 1.3, 7.4, 5.4)
picture(s, "20-prometheus-targets.png", 8.3, 1.5, 4.6, 2.4)
text_box(s, 8.3, 4.2, 4.6, 2.0, ["Target ironpass-backend", "4 of 4 UP", "Dashboard is provisioned from a ConfigMap"], size=18, color=SOFT)

# 13. troubleshooting ----------------------------------------------------
s = new_slide("Troubleshooting lab", (
    "I made two broken manifests on purpose. The first has an image tag that does not exist, so "
    "the pod goes to ImagePullBackOff. The second has a typo in the service selector, so the "
    "service has no endpoints. Describe and get endpoints show the cause, and the fix is one "
    "command each."), 13)
bullets(s, 0.6, 1.4, 4.4, 5, [
    "Broken image: ImagePullBackOff",
    "Broken service: no endpoints",
    "Tools: describe, events, endpoints",
    "Fix notes in troubleshooting/",
], size=22)
picture(s, "slides/trouble.png", 5.2, 1.4, 7.7, 5.2)

# 14. learned ------------------------------------------------------------
s = new_slide("What I learned", (
    "The biggest lesson is that each tool has one job and the pipeline joins them. A scanner "
    "that fails the build is useful, because it made me fix real findings. Small commits made "
    "problems easy to find. And cost matters, so I wrote a teardown script that removes the "
    "load balancer before the network."), 14)
bullets(s, 0.6, 1.5, 12, 5, [
    "Every tool has one job, the pipeline joins them",
    "A failing scan is useful, it made me fix real issues",
    "Infrastructure as code makes the cluster repeatable",
    "Probes, requests and metrics make an app operable",
    "Cost matters: tag everything and tear down in order",
], size=26, gap=16)

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
